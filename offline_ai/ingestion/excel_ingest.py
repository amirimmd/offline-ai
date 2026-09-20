"""Excel (.xlsx) ingestion — one row → one knowledge document."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any

from offline_ai.ingestion.base import IngestItem
from offline_ai.ingestion.json_ingest import item_from_mapping

_TEXT_KEYS = (
    "text",
    "content",
    "body",
    "full_text",
    "knowledge",
    "متن",
    "دانش",
    "متن دانش",
    "متن کامل",
    "محتوا",
)


def _norm_header(value: Any) -> str:
    return str(value or "").strip()


def _row_to_mapping(headers: list[str], values: tuple[Any, ...]) -> dict[str, Any]:
    mapping: dict[str, Any] = {}
    for i, key in enumerate(headers):
        if not key:
            continue
        val = values[i] if i < len(values) else None
        if val is None:
            continue
        if isinstance(val, str):
            val = val.strip()
            if not val:
                continue
        mapping[key] = val
    return mapping


def _fallback_text(mapping: dict[str, Any], values: tuple[Any, ...]) -> str | None:
    lower_map = {str(k).strip().lower(): v for k, v in mapping.items()}
    for key in _TEXT_KEYS:
        v = lower_map.get(key.lower())
        if v is not None and str(v).strip():
            return str(v).strip()
    parts = [str(v).strip() for v in values if v is not None and str(v).strip()]
    if not parts:
        return None
    return " | ".join(parts)


def count_excel_data_rows(path: Path) -> int:
    """Count non-header rows (approx total for progress)."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ImportError(
            "openpyxl is required for Excel ingest. Install with: pip install openpyxl"
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        max_row = int(ws.max_row or 0)
        return max(0, max_row - 1)
    finally:
        wb.close()


def iter_excel_items(
    path: Path,
    *,
    sheet: str | None = None,
    text_column: str | None = None,
) -> Iterator[IngestItem]:
    """
    Stream Excel rows as IngestItem (first row = headers).

    Prefer a text-like column (text / content / متن / دانش / …).
    Otherwise join all non-empty cells with `` | ``.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ImportError(
            "openpyxl is required for Excel ingest. Install with: pip install openpyxl"
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            return

        headers = [_norm_header(h) or f"col_{i+1}" for i, h in enumerate(header_row or ())]
        preferred = text_column.strip() if text_column else None

        for row_idx, values in enumerate(rows, start=2):
            values_t = tuple(values or ())
            if not any(v is not None and str(v).strip() for v in values_t):
                continue
            mapping = _row_to_mapping(headers, values_t)

            if preferred and preferred in mapping and str(mapping[preferred]).strip():
                text = str(mapping[preferred]).strip()
                yield IngestItem(
                    text=text,
                    source=path.name,
                    source_type="excel",
                    source_name=path.name,
                    collection_method="file",
                    metadata={"row": row_idx, **{k: v for k, v in mapping.items() if k != preferred}},
                )
                continue

            try:
                item = item_from_mapping(mapping, default_source=path.name)
                item.source_type = "excel"
                item.source_name = path.name
                item.collection_method = "file"
                item.metadata = {**(item.metadata or {}), "row": row_idx}
                yield item
                continue
            except ValueError:
                pass

            text = _fallback_text(mapping, values_t)
            if not text:
                continue
            yield IngestItem(
                text=text,
                source=path.name,
                source_type="excel",
                source_name=path.name,
                collection_method="file",
                metadata={"row": row_idx, **mapping},
            )
    finally:
        wb.close()


def load_excel_items(path: Path, **kwargs: Any) -> list[IngestItem]:
    """Materialize all Excel rows (small files). Prefer ``iter_excel_items`` for large sheets."""
    return list(iter_excel_items(path, **kwargs))
